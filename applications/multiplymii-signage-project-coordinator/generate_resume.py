import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_signage_resume():
    doc = docx.Document()

    # 0.32 in top/bottom, 0.4 in left/right for tight 1-page fit
    for section in doc.sections:
        section.top_margin = Inches(0.32)
        section.bottom_margin = Inches(0.32)
        section.left_margin = Inches(0.4)
        section.right_margin = Inches(0.4)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)

    # Color Palette
    COLOR_PRIMARY = RGBColor(15, 23, 42)      # Deep Slate #0F172A
    COLOR_ACCENT = RGBColor(2, 132, 199)      # Vivid Blue #0284C7
    COLOR_TEXT = RGBColor(51, 65, 85)         # Charcoal Slate #334155
    COLOR_MUTED = RGBColor(100, 116, 139)     # Cool Grey #64748B

    def set_font(run, name="Calibri", size_pt=8.5, bold=False, italic=False, color=COLOR_TEXT):
        run.font.name = name
        run.font.size = Pt(size_pt)
        run.bold = bold
        run.italic = italic
        run.font.color.rgb = color

    def add_section_header(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3.5)
        p.paragraph_format.space_after = Pt(1.0)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(title.upper())
        set_font(run, name="Calibri", size_pt=9.0, bold=True, color=COLOR_PRIMARY)
        
        # Bottom border
        pPr = p._p.get_or_add_pPr()
        pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="0284C7"/></w:pBdr>')
        pPr.append(pBdr)
        return p

    def add_bullet(text_parts, space_after=0.5):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.02
        p.paragraph_format.left_indent = Inches(0.16)
        for part in text_parts:
            txt = part[0]
            bld = part[1] if len(part) > 1 else False
            col = part[2] if len(part) > 2 else COLOR_TEXT
            run = p.add_run(txt)
            set_font(run, name="Calibri", size_pt=8.2, bold=bld, color=col)
        return p

    # --- HEADER ---
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(0)
    run_name = title_p.add_run("JEZREEL DAVE LEYBAG")
    set_font(run_name, name="Calibri", size_pt=15, bold=True, color=COLOR_PRIMARY)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(1.5)
    run_sub = sub_p.add_run("Signage Project Coordinator | Exterior Illuminated Signage & National Rollouts")
    set_font(run_sub, name="Calibri", size_pt=9.5, bold=True, color=COLOR_ACCENT)

    contact_p = doc.add_paragraph()
    contact_p.paragraph_format.space_before = Pt(0)
    contact_p.paragraph_format.space_after = Pt(1.0)
    
    contacts = [
        ("Email: ", True, COLOR_PRIMARY),
        ("jezreelleybag.graphics@gmail.com  |  ", False, COLOR_TEXT),
        ("Phone/WhatsApp: ", True, COLOR_PRIMARY),
        ("+63 939-399-1289  |  ", False, COLOR_TEXT),
        ("Location: ", True, COLOR_PRIMARY),
        ("Philippines (Remote CST Shift Ready: 8:00 AM – 5:00 PM CST)", False, COLOR_TEXT)
    ]
    for c_txt, c_bld, c_col in contacts:
        r = contact_p.add_run(c_txt)
        set_font(r, size_pt=8.0, bold=c_bld, color=c_col)

    links_p = doc.add_paragraph()
    links_p.paragraph_format.space_before = Pt(0)
    links_p.paragraph_format.space_after = Pt(2.5)
    
    links = [
        ("LinkedIn: ", True, COLOR_PRIMARY),
        ("linkedin.com/in/jezreel-dave-leybag-a01528152  |  ", False, COLOR_TEXT),
        ("Technical Project: ", True, COLOR_ACCENT),
        ("signquote-studio.vercel.app  |  ", False, COLOR_TEXT),
        ("Portfolio: ", True, COLOR_PRIMARY),
        ("drive.google.com/drive/folders/1_cBN8MPPfaPkyvYzinhUuSuz0kexx1tx", False, COLOR_TEXT)
    ]
    for l_txt, l_bld, l_col in links:
        r = links_p.add_run(l_txt)
        set_font(r, size_pt=8.0, bold=l_bld, color=l_col)

    pPr = links_p._p.get_or_add_pPr()
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="8" w:space="2" w:color="0F172A"/></w:pBdr>')
    pPr.append(pBdr)

    # --- PROFESSIONAL SUMMARY ---
    add_section_header("Professional Summary")
    sum_p = doc.add_paragraph()
    sum_p.paragraph_format.space_before = Pt(1.0)
    sum_p.paragraph_format.space_after = Pt(2.0)
    sum_p.paragraph_format.line_spacing = 1.05
    sum_text = [
        ("Detail-driven ", False),
        ("Signage Project Coordinator", True),
        (" with 4+ years of cross-functional operational coordination, multi-vendor quote compilation, and technical production tracking. Hands-on experience coordinating ", False),
        ("exterior illuminated signage solutions", True),
        (" (front-lit & reverse halo channel letters, raceways, lightboxes/cabinets, pylon/monuments), ", False),
        ("architectural drawing review", True),
        (" (A-sheets, wall sections, structural backing, electrical disconnects), and ", False),
        ("municipal signage code due diligence & permitting", True),
        (" across diverse jurisdictions. Creator of ", False),
        ("SignQuote Studio", True, COLOR_ACCENT),
        (" (signquote-studio.vercel.app), a dedicated commercial sign estimating and vendor RFQ workbench built on CoreBridge pricing logic, substrate yield nesting, and 50% gross margin defense. Proven track record tracking concurrent milestones across 10+ accounts, wholesale fabricators, and internal stakeholders with zero deadline slippage. Fully equipped with a Windows workstation and committed to the 8:00 AM – 5:00 PM CST schedule.", False),
    ]
    for txt, bld, *rest in sum_text:
        col = rest[0] if rest else COLOR_TEXT
        r = sum_p.add_run(txt)
        set_font(r, size_pt=8.2, bold=bld, color=col)

    # --- CORE COMPETENCIES MATRIX (COMPACT 3x2 TABLE) ---
    add_section_header("Core Competencies & Signage Capabilities")
    
    table = doc.add_table(rows=3, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    col_widths = [Inches(3.85), Inches(3.85)]
    for row in table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

    skills_data = [
        [
            ("Project & Account Coordination: ", "Multi-site rollout tracking, online PM portals, milestone delegation, vendor & client communication."),
            ("Architectural Drawing Review: ", "Reading A-sheets, elevations, wall sections, blocking/backing requirements, mounting details, electrical.")
        ],
        [
            ("Vendor Quoting & Proposals: ", "Drafting RFQs, collecting & comparing fabricator bids, spreadsheet cost consolidation, markup/margin defense."),
            ("Signage Code & Permitting: ", "Municipal zoning research, allowable sign area, setback rules, sign height, submittal packets, inspections.")
        ],
        [
            ("Signage Typologies: ", "Front-lit & halo channel letters, raceways, cabinet signs, blade/projecting signs, awnings, canopies, ADA."),
            ("Materials & Illumination: ", "Cast acrylic, ACM (3mm/4mm), polycarbonate (Lexan), aluminum returns/trim cap, LED modules, NEC Art. 600.")
        ]
    ]

    for r_idx, row_content in enumerate(skills_data):
        row = table.rows[r_idx]
        for c_idx, cell_data in enumerate(row_content):
            cell = row.cells[c_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(1.0)
            p.paragraph_format.line_spacing = 1.01
            r_title = p.add_run(cell_data[0])
            set_font(r_title, size_pt=8.0, bold=True, color=COLOR_PRIMARY)
            r_desc = p.add_run(cell_data[1])
            set_font(r_desc, size_pt=7.8, bold=False, color=COLOR_TEXT)

    # --- FEATURED TECHNICAL BENCHMARK ---
    add_section_header("Featured Technical Project & Estimating Benchmark")

    p_proj = doc.add_paragraph()
    p_proj.paragraph_format.space_before = Pt(1.5)
    p_proj.paragraph_format.space_after = Pt(0.5)
    p_proj.paragraph_format.keep_with_next = True
    r_pt1 = p_proj.add_run("SignQuote Studio — Commercial Signage Estimating & Rollout Workbench ")
    set_font(r_pt1, size_pt=8.8, bold=True, color=COLOR_PRIMARY)
    r_pt2 = p_proj.add_run("(Live Demo: signquote-studio.vercel.app)")
    set_font(r_pt2, size_pt=8.2, bold=True, color=COLOR_ACCENT)

    add_bullet([
        ("Engineered a production-ready commercial signage estimating platform modeling ", False),
        ("CoreBridge-compatible assembly pricing", True),
        (", multi-substrate yield nesting (ACM, cast acrylic, polycarbonate), machine/router labor, and strict ", False),
        ("mathematical gross margin defense", True),
        (" (COGS / (1 - GM%)) across standard commercial contract archetypes.", False)
    ])
    add_bullet([
        ("Developed automated ", False),
        ("wholesale vendor RFQ generators", True),
        (" tailored to North American fabricator standards (Direct Sign Wholesale, Gemini), specifying mandatory 1:1 paper mounting patterns, UL 48 listing, crating, and liftgate logistics.", False)
    ])
    add_bullet([
        ("Integrated AI-assisted prompt workflows to parse contractor architectural specifications and RFP emails, converting raw scope narratives into structured engineering Bills of Materials (BOM) in under 8 minutes.", False)
    ])

    # --- PROFESSIONAL EXPERIENCE ---
    add_section_header("Professional Experience")

    # ROLE 1: Morpho Studio
    p_r1 = doc.add_paragraph()
    p_r1.paragraph_format.space_before = Pt(2.0)
    p_r1.paragraph_format.space_after = Pt(0.5)
    p_r1.paragraph_format.keep_with_next = True
    
    t_role = p_r1.add_run("Project Operations & Sourcing Coordinator")
    set_font(t_role, size_pt=9.0, bold=True, color=COLOR_PRIMARY)
    t_div = p_r1.add_run(" | ")
    set_font(t_div, size_pt=9.0, bold=False, color=COLOR_MUTED)
    t_co = p_r1.add_run("Morpho Studio & Commercial Operations")
    set_font(t_co, size_pt=9.0, bold=True, color=COLOR_ACCENT)
    t_dt = p_r1.add_run(" (Jan 2024 – Present)")
    set_font(t_dt, size_pt=8.0, bold=False, color=COLOR_MUTED)

    add_bullet([
        ("Coordinated concurrent multi-vendor procurement pipelines, production deadlines, and cross-functional deliverables across 10+ operational workflows, achieving ", False),
        ("100% on-time milestone fulfillment", True),
        (" with zero material bottlenecks.", False)
    ])
    add_bullet([
        ("Audited and compared competing supplier quotations; compiled standardized pricing comparison worksheets to evaluate unit costs, lead times, freight terms, and bulk volume discounts.", False)
    ])
    add_bullet([
        ("Built and maintained centralized project tracking portals and spreadsheets in Google Workspace and Excel, logging purchase orders (POs), milestone deadlines, and vendor communication logs.", False)
    ])

    # ROLE 2: Commercial Printing & Graphics Studio
    p_r2 = doc.add_paragraph()
    p_r2.paragraph_format.space_before = Pt(2.0)
    p_r2.paragraph_format.space_after = Pt(0.5)
    p_r2.paragraph_format.keep_with_next = True

    t_role2 = p_r2.add_run("Commercial Signage & Print Production Specialist")
    set_font(t_role2, size_pt=9.0, bold=True, color=COLOR_PRIMARY)
    t_div2 = p_r2.add_run(" | ")
    set_font(t_div2, size_pt=9.0, bold=False, color=COLOR_MUTED)
    t_co2 = p_r2.add_run("Commercial Graphics & Printing Studio")
    set_font(t_co2, size_pt=9.0, bold=True, color=COLOR_ACCENT)
    t_dt2 = p_r2.add_run(" (Jan 2023 – Dec 2024)")
    set_font(t_dt2, size_pt=8.0, bold=False, color=COLOR_MUTED)

    add_bullet([
        ("Managed end-to-end signage fabrication and branded graphics projects for commercial clients, overseeing project timelines from vector art proofing through vendor handoff.", False)
    ])
    add_bullet([
        ("Reviewed architectural elevations, site layouts, and technical drawings to evaluate sign dimensions, wall substrate compatibility (masonry, drywall, siding), and structural mounting requirements.", False)
    ])
    add_bullet([
        ("Conducted preliminary code and landlord criteria due diligence; audited fabricator proofs using Adobe Acrobat Pro & Illustrator to verify scaling, PMS color matching, weep holes, and mounting patterns.", False)
    ])

    # ROLE 3: International Remote Specialist
    p_r3 = doc.add_paragraph()
    p_r3.paragraph_format.space_before = Pt(2.0)
    p_r3.paragraph_format.space_after = Pt(0.5)
    p_r3.paragraph_format.keep_with_next = True

    t_role3 = p_r3.add_run("Remote Project Coordinator & Digital Asset Specialist")
    set_font(t_role3, size_pt=9.0, bold=True, color=COLOR_PRIMARY)
    t_div3 = p_r3.add_run(" | ")
    set_font(t_div3, size_pt=9.0, bold=False, color=COLOR_MUTED)
    t_co3 = p_r3.add_run("Independent Commercial Clients (USA, Canada, UAE, Italy)")
    set_font(t_co3, size_pt=9.0, bold=True, color=COLOR_ACCENT)
    t_dt3 = p_r3.add_run(" (Dec 2020 – Dec 2024)")
    set_font(t_dt3, size_pt=8.0, bold=False, color=COLOR_MUTED)

    add_bullet([
        ("Coordinated digital production pipelines and multi-asset rollouts for international corporate accounts and consumer brands (including Fortune 500 campaigns), ensuring strict adherence to brand standards.", False)
    ])
    add_bullet([
        ("Maintained proactive written and verbal English communication across multiple US and global time zones via Microsoft Teams, email, and PM portals, facilitating seamless feedback loops and fast issue resolution.", False)
    ])

    # --- TECHNICAL SKILLS, SOFTWARE & EQUIPMENT ---
    add_section_header("Technical Skills, Software & Equipment")

    add_bullet([
        ("Project & Quoting Software: ", True),
        ("Microsoft Office (Advanced Excel with VLOOKUP/Pivots, Word, Outlook, Teams), Google Workspace (Sheets, Docs, Drive), Dropbox, CoreBridge estimating logic, Asana, Trello, ClickUp.", False)
    ])
    add_bullet([
        ("Graphic & Plan Review Tools: ", True),
        ("Adobe Acrobat Pro (PDF markup, dimensional callouts, submittal review), Adobe Illustrator (vector scaling, proof pre-flight), CorelDRAW, Bluebeam Revu (basic plan inspection).", False)
    ])
    add_bullet([
        ("Signage Technical Domain: ", True),
        ("Exterior Illuminated Channel Letters (Front-lit, Reverse Halo, Open-face), Raceways, Lightbox Cabinets, Blade Signs, Canopies, Awnings, Wayfinding/ADA, NEC Article 600 (80% power load rule), UL 48 standards.", False)
    ])
    add_bullet([
        ("Credentials & Remote Infrastructure: ", True),
        ("Google Certified AI Specialist (Google for Education & MIT RAISE - 100% Score); Windows 11 Desktop Workstation; 300 Mbps Fiber Internet + 4G Mobile Hotspot failover; Dual UPS battery backup for zero downtime.", False)
    ])

    output_path = r"d:\Jezreel Dave - Personal Branding\my-personal-website\applications\multiplymii-signage-project-coordinator\Jezreel_Dave_Leybag_Signage_Project_Coordinator.docx"
    doc.save(output_path)
    print(f"Resume generated successfully at: {output_path}")

if __name__ == "__main__":
    create_signage_resume()
